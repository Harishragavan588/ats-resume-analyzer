"""
tests/test_parser.py — Unit tests for resume_parser.py
"""

import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import tempfile
import unittest
from pathlib import Path

from services.resume_parser import parse_resume, _parse_pdf, _parse_docx


class TestResumeParser(unittest.TestCase):

    def _make_docx(self, content: str) -> Path:
        """Helper: create a real DOCX file with given text."""
        from docx import Document
        doc = Document()
        for line in content.split('\n'):
            doc.add_paragraph(line)
        tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        doc.save(tmp.name)
        tmp.close()
        return Path(tmp.name)

    def _make_pdf(self, content: str) -> Path:
        """Helper: create a minimal PDF file using reportlab."""
        from reportlab.pdfgen import canvas
        tmp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
        c = canvas.Canvas(tmp.name)
        y = 750
        for line in content.split('\n'):
            c.drawString(50, y, line)
            y -= 20
        c.save()
        tmp.close()
        return Path(tmp.name)

    # ── DOCX tests ────────────────────────────────────────

    def test_docx_extracts_text(self):
        path = self._make_docx("Jane Doe\nPython Developer\nSkills: Python, Flask, SQL")
        try:
            text = _parse_docx(path)
            self.assertIn("Jane", text)
            self.assertIn("Python", text)
        finally:
            path.unlink(missing_ok=True)

    def test_docx_empty_raises(self):
        from docx import Document
        doc = Document()
        tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        doc.save(tmp.name)
        tmp.close()
        path = Path(tmp.name)
        try:
            with self.assertRaises(RuntimeError):
                _parse_docx(path)
        finally:
            path.unlink(missing_ok=True)

    # ── PDF tests ─────────────────────────────────────────

    def test_pdf_extracts_text(self):
        path = self._make_pdf("John Smith\nSoftware Engineer\nPython Flask REST API")
        try:
            text = _parse_pdf(path)
            self.assertIn("John", text)
        finally:
            path.unlink(missing_ok=True)

    # ── Dispatch tests ────────────────────────────────────

    def test_parse_resume_unsupported_type(self):
        with tempfile.NamedTemporaryFile(suffix='.txt', delete=False) as f:
            f.write(b"hello")
            path = Path(f.name)
        try:
            with self.assertRaises(ValueError):
                parse_resume(path)
        finally:
            path.unlink(missing_ok=True)


if __name__ == '__main__':
    unittest.main()
